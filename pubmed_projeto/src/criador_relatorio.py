"""
Esse módulo é usado para criar relatórios baseados em um conjunto de resumos.
A LLM usada nessa fase de teste é o Gemini.
Atenção: esse arquivo foi feito para testes e chave de api não deve ser exposta na variável api_key
caso o arquivo com a chave se torne público.
(O módulo está incompleto ainda)
"""

from pubmed_projeto.src.buscar import procurar_resumos
from google import genai
from pubmed_projeto.src.erros import ChaveInvalidaError
from docx import Document

chave_api = "SUA_CHAVE"
modelo_gemini = "SEU_MODELO"

cliente = genai.Client(
    api_key=chave_api
)

def _validar_dados():
    """Valida a chave de API e o modelo do Gemini configurados no módulo.

    Verifica se a chave de API e o nome do modelo foram preenchidos e se o modelo
    informado está disponível na conta associada à chave.

    Args:
        None.

    Returns:
        list: Lista com os nomes de todos os modelos disponíveis para a chave de
        API configurada.

    Raises:
        ChaveInvalidaError: Se a chave de API ou o modelo ainda estiverem com os
            valores padrão, ou se o modelo informado não estiver disponível.
    """
    if chave_api == "SUA_CHAVE":
        raise ChaveInvalidaError("Você precisa colocar uma chave")

    if modelo_gemini == "SEU_MODELO":
        raise ChaveInvalidaError("Você precisa colocar um modelo")

    lista_modelos = []
    for modelo in cliente.models.list():
        lista_modelos.append(modelo.name)
    if f"models/{modelo_gemini}" not in lista_modelos:
        raise ChaveInvalidaError("Modelo inválido")

    return lista_modelos


def _gerar_resumo_inicial(resumos: dict, tema: str) -> str:
    """Gera um resumo síntese a partir dos primeiros resumos de um tema.

    Valida os dados de configuração e envia ao Gemini um prompt contendo os três
    primeiros resumos recebidos, solicitando um texto síntese sobre o tema.

    Args:
        resumos (dict): Dicionário com os resumos dos artigos. Deve conter a chave
            ``"Resumo"`` com uma lista de strings.
        tema (str): Tema sobre o qual o resumo síntese será gerado.

    Returns:
        str: Texto do resumo síntese gerado pelo modelo Gemini.

    Raises:
        ChaveInvalidaError: Se a chave de API ou o modelo configurados forem inválidos.
    """
    _validar_dados()
    primeiros_resumos = resumos["Resumo"][0:3]
    primeiros_resumos = " // ".join(primeiros_resumos)
    prompt = f"""Atue como um especialista em {tema}. Escreva um resumo síntese de 240 a 320 palavras com base EXCLUSIVAMENTE nos resumos de artigos fornecidos abaixo. Não peça mais informações, apenas gere o texto final.

    Resumos dos artigos:
    {primeiros_resumos}"""
    resumo_inicial = cliente.models.generate_content(
        model=modelo_gemini,
        contents=prompt,
    )

    return resumo_inicial.text

def gerar_relatorio(tema: str, quantidade_resumos: int):
    """Gera e salva um relatório em DOCX com resumos sobre um tema.

    Busca os resumos no PubMed, gera um resumo síntese com o Gemini e monta um
    documento contendo o resumo inicial e os resumos individuais dos artigos. O
    arquivo é salvo como ``relatorio_<tema>.docx`` no diretório atual.

    Args:
        tema (str): Tema sobre o qual o relatório será gerado.
        quantidade_resumos (int): Quantidade máxima de resumos a serem buscados.

    Returns:
        None: A função não retorna valor; o relatório é salvo em disco.

    Raises:
        ChaveInvalidaError: Se a chave de API ou o modelo configurados forem inválidos.
        Exception: Erros levantados durante a busca no PubMed ou a geração do
            documento.
    """
    document = Document()
    resumos = procurar_resumos(tema, quantidade_resumos, imprimir=False)
    resumo_inicial = _gerar_resumo_inicial(resumos, tema)
    document.add_heading(f"Relatório sobre {tema}", 0)
    document.add_heading("Resumo Inicial", level=1)
    document.add_paragraph(resumo_inicial)
    document.add_heading("Resumos", level=1)
    for numero_resumo in range(len(resumos["Resumo"])):
        document.add_heading(resumos["Título"][numero_resumo], level=2)
        document.add_paragraph(resumos["Resumo"][numero_resumo])
    document.save(f"relatorio_{tema}.docx")

if __name__ == "__main__":
    gerar_relatorio("Alzheimer", 5)
